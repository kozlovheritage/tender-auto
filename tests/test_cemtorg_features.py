"""Offline regression tests: no app import, licensing, secrets, LLM or EIS calls."""
import ast
from decimal import Decimal
import datetime
import json
import os
from pathlib import Path
import re
import tempfile
import unittest
from zipfile import ZipFile

from tendercore.auctions import auction_eligibility, auction_search_url
from tendercore.logistics import calculate, delivery_address, estimate, select_route
from tendercore.report.pdf_report import blocks, build_pdf
from tendercore.report.bundle import save_pdf, ensure_pdf, copy_report_pair
from tendercore.cemtorg_output import clean_report, current_archive

ROOT = Path(__file__).resolve().parents[1]

RESPONSE = """Решение: участвуем
--- Основная информация ---
Заказчик: ООО «Пример & контроль»
Предмет закупки: Поставка сухой строительной смеси
НМЦК/НМЦД: 1 234 567,89 руб.
НДС: включён, как указано в документах
Адрес поставки: Россия, Курская область, г. Курск, ул. Примерная, д. 1
Срок поставки: 14 календарных дней
Срок оплаты: в течение 10 рабочих дней
--- Что нужно поставить ---
ТЕКСТ_ДО_ТАБЛИЦЫ: Включая упаковку и доставку.
| № | Наименование | Артикул/Каталожный номер | Ед. изм. | Кол-во |
| --- | --- | --- | --- | --- |
| 1 | Смесь пескобетонная М300 | X-301 | т | 42,5 |
| 2 | Клей с условием A < B & C > D | — | кг | 123 |
ТЕКСТ_ПОСЛЕ_ТАБЛИЦЫ: Документ № АБ-123, п. 5.2.
--- Критичные требования ТЗ ---
Гарантия: 12 месяцев.
--- Критичные требования ТЗ ---
ПОВТОРНЫЙ_РАЗДЕЛ: Не потерять.
"""


def ast_functions(path, names, env, nested=False):
    text = path.read_text(encoding="utf-8-sig")
    tree = ast.parse(text)
    nodes = [node for node in ast.walk(tree) if isinstance(node, ast.FunctionDef) and node.name in names]
    for node in nodes:
        node.decorator_list = []
    exec(compile(ast.fix_missing_locations(ast.Module(body=nodes, type_ignores=[])), str(path), "exec"), env)
    return env


def pdf_text(path):
    import fitz
    with fitz.open(path) as doc:
        return "\n".join(page.get_text() for page in doc)


def condensed(text):
    return re.sub(r"\s+", "", text)


class AuctionTests(unittest.TestCase):
    def row(self, code, label="ФЗ 44: Электронный аукцион"):
        return {"url": f"https://zakupki.gov.ru/epz/order/notice/{code}/view/common-info.html?regNumber=0123456789012345678",
                "type": label}

    def test_confirmed_codes(self):
        for code in ("ea20", "ea15", "ea44"):
            self.assertTrue(auction_eligibility(self.row(code))[0])

    def test_not_by_open_or_auction_keyword(self):
        for code in ("ok20", "zk20", "a44", "unknown", "closed20"):
            self.assertFalse(auction_eligibility(self.row(code, "Открытый аукцион"))[0])
        self.assertFalse(auction_eligibility(self.row("ok20"))[0])

    def test_223_not_automatically_qualified(self):
        self.assertFalse(auction_eligibility(self.row("ea20", "ФЗ 223: Аукцион в электронной форме"))[0])

    def test_request_is_44_only_but_keeps_other_filters(self):
        from urllib.parse import parse_qs, urlsplit
        url = auction_search_url("https://zakupki.gov.ru/search?fz223=on&fz94=on&searchString=цемент&pageNumber=4")
        query = parse_qs(urlsplit(url).query)
        self.assertEqual(query["fz44"], ["on"])
        self.assertNotIn("fz223", query)
        self.assertNotIn("fz94", query)
        self.assertEqual(query["pageNumber"], ["4"])
        self.assertEqual(query["searchString"], ["цемент"])

    def scan_env(self, cemtorg=True, auction=True):
        today = datetime.datetime(2026, 10, 21)
        cfg = dict(cemtorg_auction_only=auction, fz44=True, fz223=True,
                   records_per_page=100)
        env = dict(cfg=cfg, CEMTORG_PROFILE=cemtorg, seen=set(), all_rows=[],
                   today=today, window_from=today - datetime.timedelta(days=6),
                   _ta_auction_eligibility=auction_eligibility)
        ast_functions(ROOT / "sampler.py", {"_accept"}, env)
        return env

    def test_scoped_filter_and_publication_window(self):
        row = dict(self.row("ok20"), reg="01", deadline=None, price=100,
                   published=datetime.datetime(2026, 10, 20))
        self.assertFalse(self.scan_env()["_accept"](row, True))
        self.assertTrue(self.scan_env(False)["_accept"](row, True))
        self.assertTrue(self.scan_env(True, False)["_accept"](row, True))
        row.update(self.row("ea20"))
        row["published"] = datetime.datetime(2026, 10, 15)
        self.assertTrue(self.scan_env()["_accept"](row, True))
        row["published"] = datetime.datetime(2026, 10, 14)
        self.assertFalse(self.scan_env()["_accept"](row, True))

    def test_all_five_pages_including_empty(self):
        class NetworkError(Exception):
            pass
        pages = []
        env = self.scan_env()
        env.update(get=lambda session, url: (pages.append(url) or type("Response", (), {"status_code": 200, "text": ""})()),
                   session=object(), set_page=lambda url, page, size: page,
                   parse_page=lambda html: [], EisNetworkError=NetworkError)
        ast_functions(ROOT / "sampler.py", {"_scan"}, env)
        env["_scan"]("url", 5, True)
        self.assertEqual(pages, [1, 2, 3, 4, 5])

    def test_profile_config_preserved(self):
        for name in ("sampler_config.txt", "profiles/cemtorg/sampler_config.txt"):
            text = (ROOT / name).read_text(encoding="utf-8-sig")
            self.assertRegex(text, r"publish_days_back:\s*6")
            self.assertRegex(text, r"cemtorg_pages_per_term:\s*5")
            self.assertRegex(text, r"cemtorg_auction_only:\s*да")


class LogisticsTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.config_path = Path(self.temp.name) / "logistics.json"
        self.config = json.loads((ROOT / "config/cemtorg_logistics.json").read_text(encoding="utf-8"))
        self.config_path.write_text(json.dumps(self.config, ensure_ascii=False), encoding="utf-8")

    def test_known_arithmetic(self):
        result = calculate(100, dict(fuel_l_per_100km=30, fuel_rub_per_l=70,
                                    other_rub_per_km=20, fixed_rub_per_trip=1000, reserve_percent=10))
        self.assertEqual(result["run"], Decimal("200"))
        self.assertEqual(result["fuel"], Decimal("4200"))
        self.assertEqual(result["total"], Decimal("10120"))

    def test_invalid_numbers_fail(self):
        for bad in (-1, "NaN", "Infinity", True, "text"):
            with self.assertRaises(ValueError):
                calculate(bad, self.config)

    def test_address_extraction(self):
        self.assertIn("Курск", delivery_address(RESPONSE))
        self.assertIsNone(delivery_address("Адрес поставки: информация отсутствует в документации"))
        self.assertEqual(delivery_address("Адрес поставки: г. Курск,\nулица Полевая\nСрок поставки: 7 дней"),
                         "г. Курск, улица Полевая")

    def test_known_scenario_and_missing_route(self):
        result = estimate(RESPONSE, self.config_path)
        self.assertIn("Синтетический", result["status"])
        self.assertIn("MAN", str(result))
        self.assertIn("Не рассчитаны", str(result))
        unknown = estimate(RESPONSE.replace("г. Курск", "г. Омск"), self.config_path)
        self.assertEqual(unknown["status"], "Не рассчитано")
        self.assertIn("Нет подтверждённого километража", str(unknown))

    def test_region_is_not_city(self):
        route, _ = select_route("Россия, Московская область, г. Серпухов", self.config["routes"])
        self.assertIsNone(route)
        route, _ = select_route("Курская область, г. Курчатов", self.config["routes"])
        self.assertIsNone(route)

    def test_multiple_addresses_and_ambiguous_routes(self):
        response = RESPONSE.replace("г. Курск, ул. Примерная, д. 1", "г. Курск; г. Москва")
        self.assertEqual(estimate(response, self.config_path)["status"], "Не рассчитано")
        route, reason = select_route("Москва и Курск", self.config["routes"])
        self.assertIsNone(route)
        self.assertIn("несколько", reason)
        response = RESPONSE.replace("г. Курск, ул. Примерная, д. 1", "г.Курск, г.Тула")
        self.assertEqual(estimate(response, self.config_path)["status"], "Не рассчитано")

    def test_absent_or_broken_config(self):
        self.config_path.write_text("{broken", encoding="utf-8")
        self.assertEqual(estimate(RESPONSE, self.config_path)["status"], "Не рассчитано")
        self.assertEqual(estimate("", self.config_path.with_name("missing"))["status"], "Не рассчитано")

    def test_bold_address_and_house_number(self):
        response = "**Место поставки товара (адрес):** г. Курск, ул. Новая, д. 12"
        self.assertIn("Синтетический", estimate(response, self.config_path)["status"])


class PDFTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.out = Path(self.temp.name)

    def test_full_source_and_repeated_headings(self):
        self.assertEqual(sum(kind == "heading" and text == "Критичные требования ТЗ"
                             for kind, text in blocks(RESPONSE)), 2)
        path = build_pdf("0123456789012345678", "30.10.2026", RESPONSE, "", self.out / "full.pdf")
        text = condensed(pdf_text(path))
        for token in ("ТЕКСТ_ДО_ТАБЛИЦЫ", "ТЕКСТ_ПОСЛЕ_ТАБЛИЦЫ", "ПОВТОРНЫЙ_РАЗДЕЛ",
                      "42,5", "123", "АБ-123", "A<B&C>D", "включён,какуказановдокументах"):
            self.assertIn(token, text)
        self.assertFalse(list(self.out.glob(".ta_pdf_*")))

    def test_long_cell_and_multi_page_table(self):
        rows = "| № | Наименование |\n|---|---|\n"
        rows += "| 1 | " + ("Длинное описание товара и условий. " * 160) + " КОНЕЦ_ЯЧЕЙКИ |\n"
        rows += "\n".join(f"| {i} | Уникальная_позиция_{i} |" for i in range(2, 80))
        path = build_pdf("0123", "01.11.2026", rows, "", self.out / "long.pdf")
        import fitz
        with fitz.open(path) as doc:
            self.assertGreater(doc.page_count, 3)
            self.assertTrue(any("DejaVuSans" in font[3] or "Arial" in font[3]
                                for font in doc[0].get_fonts()))
        text = pdf_text(path)
        self.assertIn("КОНЕЦ_ЯЧЕЙКИ", text)
        for i in range(2, 80):
            self.assertIn(f"Уникальная_позиция_{i}", text)

    def test_wide_ragged_tables_retain_extra_cells(self):
        response = "| A | B | C | D | E | F | G | H |\n|---|---|---|---|---|---|---|---|\n"
        response += "| 1 | 2 | 3 | 4 | 5 | 6 | 7 | 8 | EXTRA_CELL |"
        path = build_pdf("123", "", response, "", self.out / "wide.pdf")
        self.assertIn("EXTRA_CELL", pdf_text(path))

    def test_profile_exclusions_and_logistics(self):
        source = RESPONSE + "\n--- Возможные поставщики ---\nPRIVATE_SUPPLIER\n--- Доступность закупки в Китае: Да ---\n"
        path = save_pdf("123", "01.11.2026", source, "PRIVATE_SUPPLIER", self.out / "123.docx",
                        True, ROOT / "config/cemtorg_logistics.json")
        text = pdf_text(path)
        self.assertNotIn("PRIVATE_SUPPLIER", text)
        self.assertNotIn("Доступность закупки в Китае", text)
        self.assertIn("Предварительная оценка", text)
        self.assertIn("MAN", text)
        self.assertIn("Синтетический", text)
        other = save_pdf("124", "", RESPONSE, "GLOBAL_SUPPLIER", self.out / "124.docx")
        self.assertIn("GLOBAL_SUPPLIER", pdf_text(other))
        self.assertNotIn("Предварительная оценка логистики", pdf_text(other))

    def test_client_logo_only_on_cemtorg_first_page(self):
        import fitz
        branded = save_pdf("901", "01.11.2026", RESPONSE, "", self.out / "901.docx",
                           True, ROOT / "config/cemtorg_logistics.json")
        with fitz.open(branded) as doc:
            self.assertIn("Подготовлено для Цемторга", doc[0].get_text())
            images = doc[0].get_image_info()
            self.assertEqual(len(images), 1)
            rect = images[0]["bbox"]
            self.assertAlmostEqual((rect[2] - rect[0]) / (rect[3] - rect[1]), 625 / 209, places=3)
            self.assertTrue(all(not page.get_images() for page in list(doc)[1:]))
        generic = save_pdf("902", "", RESPONSE, "", self.out / "902.docx")
        with fitz.open(generic) as doc:
            self.assertNotIn("Подготовлено для Цемторга", doc[0].get_text())
            self.assertFalse(doc[0].get_images())

    def test_missing_requested_logo_is_explicit_error(self):
        with self.assertRaisesRegex(FileNotFoundError, "логотип Цемторга"):
            build_pdf("903", "", RESPONSE, "", self.out / "903.pdf",
                      client_logo=self.out / "missing.png")

    def test_actual_monolith_report_function_without_importing_app(self):
        from docx import Document
        from docx.shared import Pt
        env = dict(Document=Document, Pt=Pt, os=os, re=re, _FROZEN_BASE=ROOT,
                   CEMTORG_ACTIVE_TA=True, _ta_clean_report=clean_report, _ta_save_pdf=save_pdf,
                   _SERVICE_LINE_PREFIXES=())
        names = {"create_tender_report", "add_structured_report", "extract_markdown_table",
                 "add_key_value_table", "add_single_column_table", "_clean_suppliers_for_docx", "_is_service_line"}
        ast_functions(ROOT / "tender_auto.py", names, env)
        path = Path(env["create_tender_report"]("4321", "01.11.2026", RESPONSE, "", str(self.out)))
        self.assertTrue(path.is_file())
        self.assertIn("ТЕКСТ_ПОСЛЕ_ТАБЛИЦЫ", pdf_text(path.with_suffix(".pdf")))

    def test_orphan_recovery_and_copy_pair(self):
        from docx import Document
        path = self.out / "555.docx"
        doc = Document()
        doc.add_paragraph("Решение: участвуем")
        doc.add_paragraph("Адрес поставки: г. Курск")
        doc.add_table(rows=1, cols=1).cell(0, 0).text = "СОХРАНЁННАЯ_ЯЧЕЙКА"
        doc.save(path)
        ensure_pdf("555", "", path, True, ROOT / "config/cemtorg_logistics.json")
        self.assertIn("СОХРАНЁННАЯ_ЯЧЕЙКА", pdf_text(path.with_suffix(".pdf")))
        folder = self.out / "final"
        folder.mkdir()
        copy_report_pair(path, folder)
        self.assertEqual(sorted(p.suffix for p in folder.iterdir()), [".docx", ".pdf"])

    def test_archive_only_current_pairs_and_summary(self):
        folder = self.out / "final"
        date = folder / "01_11_2026"
        date.mkdir(parents=True)
        for filename in ("123.docx", "123.pdf", "OLD.pdf"):
            (date / filename).write_bytes(b"test")
        (folder / "Доступные к закупке в Китае").mkdir()
        summary = folder / "summary.xlsx"
        summary.write_bytes(b"test")
        archive = current_archive(folder, {"01.11.2026": ["temp_reports/123.docx"]}, summary)
        with ZipFile(archive) as handle:
            self.assertEqual(set(handle.namelist()), {"01_11_2026/123.docx", "01_11_2026/123.pdf", "summary.xlsx"})

    def test_summary_pdf_link_is_portable_and_real(self):
        from openpyxl import load_workbook
        self.out.joinpath("321.pdf").write_bytes(b"test")
        env = dict(os=os, Path=Path, datetime=datetime, CEMTORG_ACTIVE_TA=True)
        ast_functions(ROOT / "tender_auto.py", {"save_run_summary"}, env)
        result = ("321", "01.11.2026", str(self.out / "321.docx"), "participate",
                  False, [], "Смесь", "100", False)
        path = env["save_run_summary"]([result], str(self.out), {})
        wb = load_workbook(path)
        try:
            ws = wb.active
            self.assertEqual(ws.cell(1, ws.max_column).value, "Путь к PDF-отчёту")
            self.assertEqual(ws.cell(2, ws.max_column).hyperlink.target, "321.pdf")
            self.assertNotIn("Поставщики найдены", [cell.value for cell in ws[1]])
        finally:
            wb.close()


if __name__ == "__main__":
    unittest.main()
