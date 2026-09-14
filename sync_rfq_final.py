#!/usr/bin/env python3
"""sync_rfq_final.py — пересборка rfq_queue.json из DOCX-отчётов.

Сканирует папку output/<клиент>_<дата>/Доступные к закупке в Китае/,
читает каждый DOCX, извлекает:
  • номер тендера (из имени файла)
  • предмет закупки
  • НМЦК
  • список поставщиков с email-адресами (включая вручную вписанные)

Результат пишет в data/rfq_queue.json
"""
import json
import re
import sys
from pathlib import Path
from datetime import datetime

try:
    from docx import Document
except ImportError:
    print("❌ pip install python-docx")
    sys.exit(1)

# ─── Настройки ───────────────────────────────────────────────────────────────
OUTPUT_ROOT = Path("output")
DATA_DIR = Path("data")
QUEUE_FILE = DATA_DIR / "rfq_queue.json"

EMAIL_RE = re.compile(r"[a-zA-Z0-9._%+\-]+@[a-zA-Z0-9.\-]+\.[a-zA-Z]{2,}")
NMCK_RE = re.compile(r"(?:НМЦК|Цена|Стоимость)[:\s]*([\d\s\u00a0.,]+)\s*(?:руб|₽)", re.I)
SUBJECT_RE = re.compile(r"(?:Предмет|Объект|Наименование)[:\s]*(.+)", re.I)
SUPPLIER_SECTION_RE = re.compile(r"(?:Поставщик|Производитель|Контрагент|Партнёр|Компания)", re.I)


def find_output_dir() -> Path | None:
    """Находит самую свежую папку вида Козлов_04.09.2026 внутри output/."""
    candidates = [d for d in OUTPUT_ROOT.iterdir() if d.is_dir() and "_" in d.name]
    if not candidates:
        return None
    # Сортируем по дате в имени папки (последняя часть после _)
    def sort_key(p: Path):
        parts = p.name.rsplit("_", 1)
        try:
            return datetime.strptime(parts[-1], "%d.%m.%Y")
        except (ValueError, IndexError):
            return datetime.min
    candidates.sort(key=sort_key, reverse=True)
    return candidates[0]


def extract_info_from_docx(path: Path) -> dict:
    """Читает DOCX и извлекает данные для очереди."""
    doc = Document(str(path))
    full_text = "\n".join(p.text for p in doc.paragraphs)
    # Также читаем таблицы
    for table in doc.tables:
        for row in table.rows:
            full_text += "\n" + " | ".join(cell.text for cell in row.cells)

    tender_id = path.stem  # имя файла = номер тендера

    # НМЦК
    nmck = None
    m = NMCK_RE.search(full_text)
    if m:
        nmck_str = m.group(1).replace("\u00a0", "").replace(" ", "").replace(",", ".")
        try:
            nmck = float(nmck_str)
        except ValueError:
            pass

    # Предмет
    subject = ""
    m = SUBJECT_RE.search(full_text)
    if m:
        subject = m.group(1).strip()[:200]

    # Email-адреса поставщиков (все найденные в документе)
    emails = EMAIL_RE.findall(full_text)
    # Убираем дубли, сохраняем порядок
    seen = set()
    unique_emails = []
    for e in emails:
        e_lower = e.lower().strip()
        if e_lower not in seen:
            seen.add(e_lower)
            unique_emails.append(e_lower)

    return {
        "tender_id": tender_id,
        "subject": subject,
        "nmck": nmck,
        "suppliers": unique_emails,
        "source_report": str(path),
        "created_at": datetime.now().isoformat(timespec="seconds"),
    }


def main():
    output_dir = find_output_dir()
    if not output_dir:
        print("❌ Не найдена папка с результатами в output/")
        sys.exit(1)

    print(f"📂 Папка результатов: {output_dir}")

    # Ищем папку "Доступные к закупке в Китае"
    china_dir = None
    for d in output_dir.iterdir():
        if d.is_dir() and "Китае" in d.name:
            china_dir = d
            break

    if not china_dir:
        print("❌ Не найдена папка 'Доступные к закупке в Китае'")
        sys.exit(1)

    print(f"📂 Папка RFQ-отчётов: {china_dir}")

    # Собираем все DOCX рекурсивно
    docx_files = sorted(china_dir.rglob("*.docx"))
    if not docx_files:
        print("❌ DOCX-отчёты не найдены")
        sys.exit(1)

    print(f"📄 Найдено отчётов: {len(docx_files)}")
    print()

    queue = []
    for docx_path in docx_files:
        info = extract_info_from_docx(docx_path)
        n_suppliers = len(info["suppliers"])
        status = "✅" if n_suppliers > 0 else "⚠️"
        print(f"  {status} {info['tender_id']}: {n_suppliers} email(s) | {info['subject'][:60]}")
        queue.append(info)

    # Пишем очередь
    DATA_DIR.mkdir(parents=True, exist_ok=True)
    QUEUE_FILE.write_text(json.dumps(queue, ensure_ascii=False, indent=2), encoding="utf-8")

    print()
    print(f"✅ Очередь пересобрана: {len(queue)} тендеров → {QUEUE_FILE}")

    # Краткая сводка
    total_emails = sum(len(t["suppliers"]) for t in queue)
    no_emails = [t["tender_id"] for t in queue if not t["suppliers"]]
    print(f"   Всего email-адресов: {total_emails}")
    if no_emails:
        print(f"   ⚠️ Без поставщиков: {', '.join(no_emails)}")


if __name__ == "__main__":
    main()